"""Workbook loading, header verification and quote standardisation."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

CCYS = ["EUR", "JPY", "GBP", "CAD", "AUD", "NZD", "CHF", "NOK", "SEK"]
UNIVERSE = CCYS + ["USD"]
TENORS = ["1M", "3M"]
VOLS = ["AY1wk", "AY3M", "AY6M", "JPY3M", "AUD3M", "VIX"]

# Market quote convention of each spot column in the workbook.
QUOTE = {
    "EUR": "USD per EUR", "JPY": "JPY per USD", "GBP": "USD per GBP",
    "CAD": "CAD per USD", "AUD": "USD per AUD", "NZD": "USD per NZD",
    "CHF": "CHF per USD", "NOK": "NOK per USD", "SEK": "SEK per USD",
}
INVERTED = [c for c in CCYS if QUOTE[c].endswith("per USD")]

DAYCOUNT = 360.0  # ACT/360 money-market accrual (assumption, applied to all currencies)

VOL_DESCRIPTION = {
    "AY1wk": "AUD/JPY option implied vol, 1-week tenor",
    "AY3M": "AUD/JPY option implied vol, 3-month tenor",
    "AY6M": "AUD/JPY option implied vol, 6-month tenor",
    "JPY3M": "USD/JPY option implied vol, 3-month tenor",
    "AUD3M": "AUD/USD option implied vol, 3-month tenor",
    "VIX": "CBOE VIX index (30-day S&P 500 implied vol)",
}


def expected_layout() -> dict[str, str]:
    """Column letter -> header text, as described in the assignment PDF."""
    lay = {"A": "Date"}
    for j, c in enumerate(CCYS):
        lay[get_column_letter(2 + j)] = c                    # B-J
        lay[get_column_letter(12 + j)] = f"{c} ROR"          # L-T
    lay["U"] = "USD ROR"
    for j, c in enumerate(UNIVERSE):
        lay[get_column_letter(23 + j)] = f"{c} 1M"           # W-AF
        lay[get_column_letter(33 + j)] = f"{c} 3M"           # AG-AP
    for j, v in enumerate(VOLS):
        lay[get_column_letter(44 + j)] = v                   # AR-AW
    return lay


@dataclass
class Dataset:
    dates: pd.DatetimeIndex
    quotes: pd.DataFrame            # raw spot quotes, market convention
    usd_per_ccy: pd.DataFrame       # X_{i,t}: USD per 1 unit of currency i
    provided_ror: pd.DataFrame      # workbook columns L-U
    rates: dict[str, pd.DataFrame]  # tenor -> annualised decimal rates, columns UNIVERSE
    vols_raw: pd.DataFrame          # vol levels in vol points, NaN where missing
    audit: dict = field(default_factory=dict)


def _find_header_row(raw: pd.DataFrame) -> int:
    for i in range(min(len(raw), 50)):
        if str(raw.iat[i, 0]).strip().lower() == "date":
            return i
    raise ValueError("No 'Date' header found in first column")


def load_dataset(path: str | Path) -> Dataset:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Required workbook missing: {path}")
    sheets = pd.read_excel(path, sheet_name=None, header=None)
    audit: dict = {"file": path.name, "sheets": {k: v.shape for k, v in sheets.items()}}
    candidates = [k for k, v in sheets.items() if len(v) and (v.iloc[:50, 0].astype(str).str.lower() == "date").any()]
    if len(candidates) != 1:
        raise ValueError(f"Expected one data sheet with a Date header, found {candidates}")
    sheet = candidates[0]
    raw = sheets[sheet]
    h = _find_header_row(raw)
    headers = [None if pd.isna(x) else str(x).strip() for x in raw.iloc[h]]
    audit["sheet"] = sheet
    audit["header_row_excel"] = h + 1

    layout = expected_layout()
    mismatches = {}
    for letter, name in layout.items():
        col = _col_index(letter)
        got = headers[col] if col < len(headers) else None
        if got != name:
            mismatches[letter] = (name, got)
    if mismatches:
        raise ValueError(f"Workbook layout differs from assignment description: {mismatches}")
    audit["layout_verified"] = True

    body = raw.iloc[h + 1:].copy()
    body.columns = [hd if hd is not None else f"_blank{j}" for j, hd in enumerate(headers)]
    body = body.loc[:, [c for c in body.columns if not c.startswith("_blank")]]
    body = body.dropna(how="all")
    body["Date"] = pd.to_datetime(body["Date"])
    n_before = len(body)
    audit["input_sorted"] = bool(body["Date"].is_monotonic_increasing)
    body = body.sort_values("Date", kind="stable")
    dup = body["Date"].duplicated(keep="last")
    audit["duplicate_dates_dropped"] = int(dup.sum())
    body = body.loc[~dup].set_index("Date")
    body = body.apply(pd.to_numeric, errors="coerce")
    audit["rows"] = len(body)
    audit["rows_before_dedup"] = n_before

    quotes = body[CCYS].astype(float)
    if quotes.isna().any().any():
        raise ValueError("Missing spot quotes; no fill policy defined for spot data")
    usd_per = quotes.copy()
    usd_per[INVERTED] = 1.0 / quotes[INVERTED]

    rates = {}
    for t in TENORS:
        r = body[[f"{c} {t}" for c in UNIVERSE]].astype(float) / 100.0
        r.columns = UNIVERSE
        if r.isna().any().any():
            raise ValueError(f"Missing {t} rates")
        rates[t] = r

    vols = body[VOLS].astype(float)
    provided = body[[f"{c} ROR" for c in CCYS] + ["USD ROR"]].astype(float)
    provided.columns = UNIVERSE

    idx = body.index
    audit["start"], audit["end"] = idx[0], idx[-1]
    audit["weekday_counts"] = pd.Series(idx.dayofweek).value_counts().sort_index().to_dict()
    audit["calendar_gaps"] = pd.Series(np.diff(idx.values).astype("timedelta64[D]").astype(int)).value_counts().sort_index().to_dict()
    audit["missing"] = body.isna().sum()[lambda s: s > 0].to_dict()
    audit["vix_missing_dates"] = list(vols.index[vols["VIX"].isna()])
    return Dataset(idx, quotes, usd_per, provided, rates, vols, audit)


def _col_index(letter: str) -> int:
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def filled_vols(vols_raw: pd.DataFrame) -> pd.DataFrame:
    """Carry the last *observed* value forward (holidays); never fill before a series starts."""
    out = vols_raw.copy()
    for c in out:
        first = out[c].first_valid_index()
        out.loc[first:, c] = out.loc[first:, c].ffill()
    return out


def staleness_by_year(vols_raw: pd.DataFrame) -> pd.DataFrame:
    """Share of days on which the quoted level is unchanged from the previous observation."""
    same = vols_raw.ffill().diff().eq(0) & vols_raw.notna()
    return same.groupby(vols_raw.index.year).mean()
